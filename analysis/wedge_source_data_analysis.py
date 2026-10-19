"""Audit the local 2026 PIEZO1 wedge Source Data without merging it into Duke.

Source: Data/41467_2026_76927_MOESM4_ESM.xlsx, sheets Fig3A,
FigS4A, Fig1J, Fig1K. Workbook columns preserve identity only within a
condition unless explicitly labelled otherwise. Reported summaries count cells.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import least_squares
from scipy.stats import mannwhitneyu, wilcoxon


HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent / "Data" / "41467_2026_76927_MOESM4_ESM.xlsx"
INTERVAL_COLUMNS = {50: (2, 10), 100: (11, 17), 200: (18, 25),
                    500: (26, 31), 1000: (32, 39)}


def numeric(value) -> bool:
    return isinstance(value, (int, float)) and np.isfinite(value)


def write_csv(name: str, rows: list[dict]) -> None:
    columns = list(dict.fromkeys(key for row in rows for key in row))
    with (HERE / name).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def recovery_data(workbook) -> list[dict]:
    data = []
    s = workbook["FigS4A"]
    for interval, (first, last) in INTERVAL_COLUMNS.items():
        for col in range(first, last + 1):
            for pulse in range(1, 7):
                value = s.cell(pulse + 2, col).value
                if numeric(value):
                    data.append({"source_sheet": "FigS4A", "workbook_column": col,
                                 "interval_ms": interval, "pulse_number": pulse,
                                 "ratio_to_first": float(value),
                                 "near_zero_after_first": int(pulse > 1 and value < 1e-6)})
    # Fig3A repeats the WT 50 ms group. Check the source duplication exactly;
    # do not count the same WT patch twice in the output.
    fig3 = workbook["Fig3A"]
    for col in range(2, 11):
        for pulse in range(1, 7):
            a = fig3.cell(pulse + 2, col).value
            b = s.cell(pulse + 2, col).value
            assert a == b, (col, pulse, a, b)
    return data


def mutant_data(workbook) -> list[dict]:
    s = workbook["Fig3A"]
    return [{"source_sheet": "Fig3A", "workbook_column": col,
             "pulse_number": pulse, "ratio_to_first": float(s.cell(pulse + 2, col).value)}
            for col in range(11, 24) for pulse in range(1, 7)
            if numeric(s.cell(pulse + 2, col).value)]


def afm_data(workbook) -> list[dict]:
    result = []
    tau, remaining = workbook["Fig1J"], workbook["Fig1K"]
    for row in range(4, 10):
        speed = float(tau.cell(row, 1).value)
        assert speed == float(remaining.cell(row, 1).value)
        for col in range(2, 9):
            t, f = tau.cell(row, col).value, remaining.cell(row, col).value
            if numeric(t) and numeric(f):
                result.append({"source_sheets": "Fig1J;Fig1K",
                               "cell_label": f"Cell{col-1}",
                               "speed_um_per_ms": speed,
                               "force_relaxation_tau_ms": float(t),
                               "force_remaining_pct": float(f)})
    return result


def fit_second_pulse(intervals: np.ndarray, ratios: np.ndarray) -> dict:
    fit = least_squares(lambda p: 1 - p[0] * np.exp(-intervals / p[1]) - ratios,
                        x0=[.2, 500.], bounds=([0., 1.], [1., 100000.]))
    return {"depletion_fraction": float(fit.x[0]), "tau_ms": float(fit.x[1]),
            "rmse_ratio": float(np.sqrt(np.mean(fit.fun ** 2)))}


def main() -> None:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    wt = recovery_data(workbook)
    mutant = mutant_data(workbook)
    afm = afm_data(workbook)
    write_csv("wedge_recovery_wt.csv", wt)
    write_csv("wedge_recovery_mutant_50ms.csv", mutant)
    write_csv("wedge_afm_velocity.csv", afm)

    interval_summary = []
    for interval in INTERVAL_COLUMNS:
        values = [r for r in wt if r["interval_ms"] == interval]
        cells = {r["workbook_column"] for r in values}
        for pulse in range(1, 7):
            x = [r["ratio_to_first"] for r in values if r["pulse_number"] == pulse]
            interval_summary.append({"interval_ms": interval, "pulse_number": pulse,
                                     "cells_in_group": len(cells), "observed": len(x),
                                     "median_ratio": float(np.median(x)),
                                     "mean_ratio": float(np.mean(x)),
                                     "q25_ratio": float(np.quantile(x, .25)),
                                     "q75_ratio": float(np.quantile(x, .75)),
                                     "near_zero_count": sum(y < 1e-6 for y in x)})
    write_csv("wedge_recovery_interval_summary.csv", interval_summary)
    afm_summary = []
    for speed in sorted({r["speed_um_per_ms"] for r in afm}):
        rows = [r for r in afm if r["speed_um_per_ms"] == speed]
        afm_summary.append({"speed_um_per_ms": speed, "cells": len(rows),
                            "mean_tau_ms": float(np.mean([r["force_relaxation_tau_ms"] for r in rows])),
                            "median_tau_ms": float(np.median([r["force_relaxation_tau_ms"] for r in rows])),
                            "mean_force_remaining_pct": float(np.mean([r["force_remaining_pct"] for r in rows])),
                            "median_force_remaining_pct": float(np.median([r["force_remaining_pct"] for r in rows]))})
    write_csv("wedge_afm_velocity_summary.csv", afm_summary)

    sixth_wt = [r["ratio_to_first"] for r in wt if r["interval_ms"] == 50 and r["pulse_number"] == 6]
    sixth_mut = [r["ratio_to_first"] for r in mutant if r["pulse_number"] == 6]
    intervals = np.array(sorted(INTERVAL_COLUMNS), dtype=float)
    second = np.array([next(r["median_ratio"] for r in interval_summary if
                            r["interval_ms"] == interval and r["pulse_number"] == 2)
                       for interval in intervals])
    full_fit = fit_second_pulse(intervals, second)
    jackknife = [{"omitted_interval_ms": int(intervals[i]),
                  **fit_second_pulse(np.delete(intervals, i), np.delete(second, i))}
                 for i in range(len(intervals))]
    write_csv("wedge_recovery_tau_sensitivity.csv", jackknife)
    paired = {}
    for key in ("force_relaxation_tau_ms", "force_remaining_pct"):
        fast = {r["cell_label"]: r[key] for r in afm if r["speed_um_per_ms"] == .1}
        slow = {r["cell_label"]: r[key] for r in afm if r["speed_um_per_ms"] == .0025}
        cells = sorted(fast.keys() & slow.keys())
        paired[key] = {"n": len(cells), "slow_median": float(np.median([slow[c] for c in cells])),
                       "fast_median": float(np.median([fast[c] for c in cells])),
                       "one_sided_wilcoxon_p": float(wilcoxon([slow[c] for c in cells],
                                                               [fast[c] for c in cells],
                                                               alternative="greater").pvalue)}
    summary = {"source": SOURCE.name,
               "recovery": {"WT_50ms_cells": len(sixth_wt),
                            "wedgeDEL_50ms_cells": len(sixth_mut),
                            "WT_sixth_median": float(np.median(sixth_wt)),
                            "wedgeDEL_sixth_median": float(np.median(sixth_mut)),
                            "sixth_pulse_Mann_Whitney_two_sided_p":
                                float(mannwhitneyu(sixth_wt, sixth_mut).pvalue),
                            "WT_interval_group_cell_counts":
                                {str(dt): len({r["workbook_column"] for r in wt
                                               if r["interval_ms"] == dt}) for dt in INTERVAL_COLUMNS},
                            "near_zero_1000ms_after_first": sum(r["near_zero_after_first"]
                                                              for r in wt if r["interval_ms"] == 1000),
                            "second_pulse_median_fit_diagnostic_only": full_fit,
                            "tau_leave_one_interval_out_ms":
                                [r["tau_ms"] for r in jackknife]},
               "afm": {"cell_count": 7, "complete_six_speed_cells": 6,
                       "paired_fast_slow": paired}}
    (HERE / "wedge_source_data_results.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    for interval in INTERVAL_COLUMNS:
        data = [r for r in interval_summary if r["interval_ms"] == interval]
        axes[0].plot([r["pulse_number"] for r in data],
                     [r["median_ratio"] for r in data], marker="o", label=f"{interval} ms")
    axes[0].plot(range(1, 7), [np.median([r["ratio_to_first"] for r in mutant
                                        if r["pulse_number"] == pulse])
                               for pulse in range(1, 7)], color="black", linestyle="--",
                 marker="s", label="wedge Δ, 50 ms")
    axes[0].set(xlabel="Pulse number", ylabel="Peak current / first peak",
                title="Repeated pressure pulses")
    axes[0].legend(fontsize=8)
    for cell in sorted({r["cell_label"] for r in afm}):
        rows = sorted([r for r in afm if r["cell_label"] == cell],
                      key=lambda r: r["speed_um_per_ms"])
        axes[1].plot([r["speed_um_per_ms"] for r in rows],
                     [r["force_relaxation_tau_ms"] for r in rows],
                     color="gray", alpha=.45, linewidth=1)
    axes[1].plot([r["speed_um_per_ms"] for r in afm_summary],
                 [r["median_tau_ms"] for r in afm_summary],
                 marker="o", color="#126ea0", linewidth=2, label="Cell median")
    axes[1].set(xscale="log", xlabel="AFM indentation speed (μm/ms)",
                ylabel="AFM force-relaxation τ (ms)",
                title="Mechanical relaxation, not current decay")
    axes[1].legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(HERE / "wedge_recovery_velocity.png", dpi=180)
    plt.close(fig)
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
