"""Independent audit of the 2026 PIEZO2 FLNB perturbation source data.

Read-only extraction from the Nature figure-source workbooks. Cell observations
are not paired across conditions or endpoints; Fig. 3j is a group mean curve.
No Duke data or model fitting are used here.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.stats import mannwhitneyu


HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "Data"
FILES = {"fig2": "41586_2026_10182_MOESM5_ESM.xlsx",
         "fig3": "41586_2026_10182_MOESM6_ESM.xlsx",
         "ed8": "41586_2026_10182_MOESM11_ESM.xlsx"}
GROUPS = [
    ("Piezo2", "indentation_threshold_um", "fig3", "Fig3k", 42),
    ("Piezo2_FLNB_KO", "indentation_threshold_um", "fig3", "Fig3k", 43),
    ("Piezo2_delta_IDR5", "indentation_threshold_um", "fig3", "Fig3k", 44),
    ("Piezo2", "peak_current_density_pA_per_pF", "ed8", "ED8a", 1),
    ("Piezo2_FLNB_KO", "peak_current_density_pA_per_pF", "ed8", "ED8a", 2),
    ("Piezo2_delta_IDR5", "peak_current_density_pA_per_pF", "ed8", "ED8a", 3),
    ("Piezo1", "indentation_threshold_um", "ed8", "ED8b", 5),
    ("Piezo1_FLNB_KO", "indentation_threshold_um", "ed8", "ED8b", 6),
    ("Piezo1", "peak_current_density_pA_per_pF", "ed8", "ED8c", 8),
    ("Piezo1_FLNB_KO", "peak_current_density_pA_per_pF", "ed8", "ED8c", 9),
    ("Untransfected", "swelling_current_density_pA_per_pF", "fig2", "Fig2c", 5),
    ("Piezo2_CytoD", "swelling_current_density_pA_per_pF", "fig2", "Fig2c", 6),
    ("Piezo2", "swelling_current_density_pA_per_pF", "fig2", "Fig2c", 7),
]


def is_number(value) -> bool:
    return isinstance(value, (int, float)) and np.isfinite(value)


def write_csv(name: str, rows: list[dict]) -> None:
    with (HERE / name).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def extract_cells(books: dict) -> list[dict]:
    rows = []
    for condition, endpoint, book, panel, column in GROUPS:
        sheet = books[book]["Figure Data"]
        # ED8 includes thousands of unrelated MINFLUX rows farther down. The
        # electrophysiology blocks used here end before row 80.
        for row, (value,) in enumerate(sheet.iter_rows(min_row=3, max_row=80,
                                                        min_col=column, max_col=column,
                                                        values_only=True), start=3):
            if is_number(value):
                rows.append({"source_file": FILES[book], "figure_panel": panel,
                             "source_row": row, "source_column": column,
                             "condition": condition, "endpoint": endpoint,
                             "value": float(value)})
    return rows


def extract_curve(sheet) -> list[dict]:
    rows = []
    for row in range(4, 25):
        depth = sheet.cell(row, 31).value
        if not is_number(depth):
            continue
        for condition, start in (("Piezo2", 32), ("Piezo2_delta_IDR5", 35),
                                 ("Piezo2_FLNB_KO", 38)):
            mean, sem, count = [sheet.cell(row, col).value for col in range(start, start+3)]
            if all(is_number(v) for v in (mean, sem, count)):
                rows.append({"source_file": FILES["fig3"], "figure_panel": "Fig3j",
                             "indentation_um": float(depth), "condition": condition,
                             "mean_current_density_pA_per_pF": float(mean),
                             "sem_pA_per_pF": float(sem), "cell_count_at_depth": int(count)})
    return rows


def group_values(rows: list[dict], endpoint: str, condition: str) -> np.ndarray:
    return np.array([r["value"] for r in rows if r["endpoint"] == endpoint and
                     r["condition"] == condition], dtype=float)


def compare(a: np.ndarray, b: np.ndarray, rng: np.random.Generator) -> dict:
    aa = rng.choice(a, size=(20000, len(a)), replace=True).mean(axis=1)
    bb = rng.choice(b, size=(20000, len(b)), replace=True).mean(axis=1)
    return {"control_n": len(a), "perturbation_n": len(b),
            "control_mean": float(np.mean(a)), "perturbation_mean": float(np.mean(b)),
            "control_median": float(np.median(a)), "perturbation_median": float(np.median(b)),
            "mean_difference": float(np.mean(b)-np.mean(a)),
            "bootstrap_95pct_mean_difference": [float(v) for v in np.quantile(bb-aa, [.025, .975])],
            "mean_ratio": float(np.mean(b)/np.mean(a)),
            "bootstrap_95pct_mean_ratio": [float(v) for v in np.quantile(bb/aa, [.025, .975])],
            "Mann_Whitney_two_sided_p": float(mannwhitneyu(a, b, alternative="two-sided").pvalue)}


def main() -> None:
    books = {key: load_workbook(DATA / filename, read_only=True, data_only=True)
             for key, filename in FILES.items()}
    cells = extract_cells(books)
    curve = extract_curve(books["fig3"]["Figure Data"])
    write_csv("piezo2_flnb_cell_values.csv", cells)
    write_csv("piezo2_flnb_indentation_curve.csv", curve)
    rng = np.random.default_rng(27)
    comparisons = {}
    for channel in ("Piezo2", "Piezo1"):
        for endpoint in ("indentation_threshold_um", "peak_current_density_pA_per_pF"):
            comparisons[f"{channel}_{endpoint}"] = compare(
                group_values(cells, endpoint, channel),
                group_values(cells, endpoint, f"{channel}_FLNB_KO"), rng)
    swelling = {condition: {"n": int(len(group_values(cells,
                    "swelling_current_density_pA_per_pF", condition))),
                "median": float(np.median(group_values(cells,
                    "swelling_current_density_pA_per_pF", condition))),
                "mean": float(np.mean(group_values(cells,
                    "swelling_current_density_pA_per_pF", condition)))}
                for condition in ("Untransfected", "Piezo2", "Piezo2_CytoD")}
    results = {"comparisons": comparisons, "swelling_CytoD": swelling,
               "limitations": "Different conditions are independent; threshold and peak are not cell-paired; Fig3j contains group means with varying cell counts."}
    (HERE / "piezo2_flnb_results.json").write_text(
        json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")

    fig, axes = plt.subplots(1, 3, figsize=(13, 4))
    colors = {"Piezo2": "#196da0", "Piezo2_FLNB_KO": "#dd8335",
              "Piezo2_delta_IDR5": "#7c5da6"}
    for condition in colors:
        data = [r for r in curve if r["condition"] == condition]
        x = np.array([r["indentation_um"] for r in data])
        y = np.array([r["mean_current_density_pA_per_pF"] for r in data])
        e = np.array([r["sem_pA_per_pF"] for r in data])
        axes[0].plot(x, y, label=condition.replace("_", " "), color=colors[condition])
        axes[0].fill_between(x, y-e, y+e, color=colors[condition], alpha=.13)
    axes[0].set(xlabel="Glass-probe indentation (μm)",
                ylabel="Mean current density (pA/pF)", title="Fig. 3j: grouped curve")
    axes[0].legend(fontsize=7)
    for ax, endpoint, title, ylabel in (
        (axes[1], "indentation_threshold_um", "Cell threshold", "Indentation (μm)"),
        (axes[2], "peak_current_density_pA_per_pF", "Cell peak current", "Current density (pA/pF)")):
        groups = [("Piezo2", "P2"), ("Piezo2_FLNB_KO", "P2 FLNB KO"),
                  ("Piezo1", "P1"), ("Piezo1_FLNB_KO", "P1 FLNB KO")]
        for i, (condition, _) in enumerate(groups):
            values = group_values(cells, endpoint, condition)
            ax.scatter(np.full(len(values), i) + rng.uniform(-.1, .1, len(values)),
                       values, s=14, alpha=.45,
                       color="#196da0" if "KO" not in condition else "#dd8335")
            ax.plot([i-.2, i+.2], [np.median(values)]*2, color="black", linewidth=2)
        ax.set_xticks(range(len(groups)), [label for _, label in groups], rotation=25)
        ax.set(ylabel=ylabel, title=title)
    fig.tight_layout()
    fig.savefig(HERE / "piezo2_flnb_perturbation.png", dpi=180)
    plt.close(fig)
    print(json.dumps(results, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
