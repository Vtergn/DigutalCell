"""Read-only inventory of the PIEZO2 OSF electrophysiology source workbook."""

from __future__ import annotations

import json
from pathlib import Path

import requests
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
SOURCE = "https://api.osf.io/v2/nodes/2cgxn/files/osfstorage/"
LOCAL = HERE.parent / "Data" / "2026_ncb_piezo2" / "Electrophysiology Source Data.xlsx"


def main() -> None:
    files = requests.get(SOURCE, timeout=30).json()["data"]
    listed = [{"name": item["attributes"]["name"],
               "bytes": item["attributes"].get("size"),
               "download": item.get("links", {}).get("download")}
              for item in files]
    if not LOCAL.exists():
        raise FileNotFoundError("Download the public OSF source workbook before auditing")
    workbook = load_workbook(LOCAL, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook:
        headings = []
        for row in sheet.iter_rows(min_row=1, max_row=min(3, sheet.max_row),
                                   values_only=True):
            headings.append([{"column": i+1, "value": str(value)}
                             for i, value in enumerate(row) if value is not None][:12])
        sheets.append({"name": sheet.title, "rows": sheet.max_row,
                       "columns": sheet.max_column, "first_three_rows": headings})
    result = {"osf_api": SOURCE, "listed_files": listed,
              "local_bytes": LOCAL.stat().st_size, "sheets": sheets,
              "grain_audit": "Figure 1 holds normalized-current versus pressure columns, not timestamped patch-clamp samples. Figure 2 holds all-point current histograms; Figure 6 holds per-cell reversal potentials. The project lists one source workbook, not raw 10-kHz recordings."}
    (HERE / "osf_2cgxn_audit.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"listed_files": listed,
                      "sheets": [(s["name"], s["rows"], s["columns"]) for s in sheets]},
                     ensure_ascii=False))


if __name__ == "__main__":
    main()
